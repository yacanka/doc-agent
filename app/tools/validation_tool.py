"""Validation helpers for generated document outputs."""
from __future__ import annotations

import json
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from app.storage.db import utc_now

_OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx"}
_VOLATILE_ENTRIES = {"docProps/core.xml"}
_EXPECTED_PREFIXES = {
    ".docx": ("word/", "_rels/", "docProps/", "[Content_Types].xml"),
    ".xlsx": ("xl/", "_rels/", "docProps/", "[Content_Types].xml"),
    ".pptx": ("ppt/", "_rels/", "docProps/", "[Content_Types].xml"),
}


@dataclass(frozen=True)
class ValidationResult:
    """Structured outcome for an output document validation."""

    valid: bool
    high_risk: bool
    parser: str
    output_path: str
    report_path: str
    added_entries: list[str]
    removed_entries: list[str]
    changed_entries: list[str]
    warnings: list[str]
    created_at: str


def validate_output_file(original_path: Path, output_path: Path, workspace_root: Path) -> ValidationResult:
    """Validate parser readability and Office package changes, then write JSON report."""
    warnings = _parser_warnings(output_path)
    package_changes = _compare_office_packages(original_path, output_path)
    high_risk = bool(warnings or _unexpected_structure(output_path, package_changes))
    report_path = _write_change_report(workspace_root, output_path, warnings, package_changes, high_risk)
    return ValidationResult(
        valid=not warnings,
        high_risk=high_risk,
        parser=_parser_name(output_path),
        output_path=str(output_path),
        report_path=str(report_path),
        warnings=warnings,
        created_at=utc_now(),
        **package_changes,
    )


def validate_replacements(path: Path, replacements: dict[str, str], changed_count: int) -> dict[str, Any]:
    """Validate a replacement by confirming expected visible text changes."""
    from app.tools.document_tool import extract_text

    text = extract_text(path)
    missing = [value for value in replacements.values() if value not in text]
    leftovers = [key for key in replacements if key in text]
    return {
        "created_at": utc_now(),
        "document": str(path),
        "changed_count": changed_count,
        "missing_replacements": missing,
        "leftover_sources": leftovers,
        "valid": not missing and not leftovers and changed_count > 0,
    }


def write_report(report: dict[str, Any], output_path: Path) -> Path:
    """Write a validation report next to the generated document."""
    report_path = output_path.with_suffix(output_path.suffix + ".validation.json")
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report_path


def _parser_warnings(output_path: Path) -> list[str]:
    try:
        _open_with_parser(output_path)
        return []
    except Exception as error:
        return [f"Parser could not open output: {error}"]


def _open_with_parser(path: Path) -> None:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        Document(path)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        load_workbook(path, read_only=True, data_only=False).close()
    elif suffix == ".pptx":
        from pptx import Presentation

        Presentation(path)
    elif suffix == ".pdf":
        from app.tools.pdf_tool import extract_text

        extract_text(path)
    else:
        path.open("rb").close()


def _compare_office_packages(original_path: Path, output_path: Path) -> dict[str, list[str]]:
    if not _is_office_file(original_path) or not _is_office_file(output_path):
        return {"added_entries": [], "removed_entries": [], "changed_entries": []}
    before = _zip_manifest(original_path)
    after = _zip_manifest(output_path)
    return {
        "added_entries": sorted(set(after) - set(before)),
        "removed_entries": sorted(set(before) - set(after)),
        "changed_entries": sorted(_changed_entries(before, after)),
    }


def _zip_manifest(path: Path) -> dict[str, tuple[int, int]]:
    with zipfile.ZipFile(path) as package:
        return {info.filename: (info.CRC, info.file_size) for info in package.infolist()}


def _changed_entries(before: dict[str, tuple[int, int]], after: dict[str, tuple[int, int]]) -> list[str]:
    shared = set(before).intersection(after)
    return [name for name in shared if before[name] != after[name]]


def _unexpected_structure(output_path: Path, changes: dict[str, list[str]]) -> bool:
    if output_path.suffix.lower() not in _OFFICE_EXTENSIONS:
        return False
    structural_changes = changes["added_entries"] + changes["removed_entries"]
    return any(_is_structural_entry(output_path.suffix.lower(), name) for name in structural_changes)


def _is_structural_entry(suffix: str, entry_name: str) -> bool:
    if entry_name in _VOLATILE_ENTRIES:
        return False
    return not entry_name.startswith(_EXPECTED_PREFIXES[suffix]) or entry_name.endswith(".rels")


def _write_change_report(
    workspace_root: Path,
    output_path: Path,
    warnings: list[str],
    changes: dict[str, list[str]],
    high_risk: bool,
) -> Path:
    report_path = _report_path(workspace_root, output_path)
    result = {
        "created_at": utc_now(),
        "output_path": str(output_path),
        "parser": _parser_name(output_path),
        "valid": not warnings,
        "high_risk": high_risk,
        "warnings": warnings,
        **changes,
    }
    report_path.write_text(json.dumps(result, indent=2), encoding="utf-8")
    return report_path


def _report_path(workspace_root: Path, output_path: Path) -> Path:
    outputs_dir = workspace_root / "outputs"
    outputs_dir.mkdir(parents=True, exist_ok=True)
    name = f"{output_path.stem}.validation.json"
    return outputs_dir / name


def _parser_name(path: Path) -> str:
    return {".docx": "python-docx", ".xlsx": "openpyxl", ".pptx": "python-pptx", ".pdf": "pymupdf/pypdf"}.get(
        path.suffix.lower(), "binary"
    )


def _is_office_file(path: Path) -> bool:
    return path.suffix.lower() in _OFFICE_EXTENSIONS and zipfile.is_zipfile(path)


def result_to_dict(result: ValidationResult) -> dict[str, Any]:
    """Convert a validation result dataclass into a JSON-safe dictionary."""
    return asdict(result)
